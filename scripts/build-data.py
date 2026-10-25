#!/usr/bin/env python3
"""
Builds src/data/units.json from the private contributions spreadsheet.

The spreadsheet itself contains personal data and must NEVER be committed.
Put it into ./private/ (git-ignored) and run:

    pip install openpyxl
    python scripts/build-data.py private/"ЖК Компаньйон - внески.xlsx" --as-of 2026-10-20

Only anonymous data is exported: building / section / floor / unit type /
unit numbers / total amount received. No names, phones, e-mails, comments.
"""
import argparse
import datetime as dt
import json
import pathlib
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

SUMMARY_SHEET = "Підсумок"
PAYMENT_SHEETS = ["Сходи", "Шлагбауми", "Решта"]
UNIT_TYPES = {"квартира": "flat", "апартамент": "apt"}  # комори are skipped
ADDRESS = {1: "31", 2: "33", 3: "35"}
# Placeholder numbers in the source -> number shown on the site
NUMBER_OVERRIDES = {(3, "1000"): "116", (3, "1001"): "117", (3, "1002"): "118"}


def norm(v):
    """Excel gives 12.0 for 12 — turn numbers into clean strings."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default="src/data/units.json")
    ap.add_argument("--as-of", help="YYYY-MM-DD; default = latest payment date in the file")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[SUMMARY_SHEET]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}

    units, totals = [], {"stairs": 0, "barriers": 0, "rest": 0}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[col["Будинок"]] is None:
            continue
        kind = UNIT_TYPES.get(row[col["Тип"]])
        if not kind:
            continue
        b = int(row[col["Будинок"]])
        number = norm(row[col["Новий номер"]])
        number = NUMBER_OVERRIDES.get((b, number), number)
        stairs = int(row[col["Сходи"]] or 0)
        barriers = int(row[col["Шлагбауми"]] or 0)
        rest = int(row[col["Решта"]] or 0)
        totals["stairs"] += stairs
        totals["barriers"] += barriers
        totals["rest"] += rest
        u = {
            "b": b,
            "s": int(row[col["Секція"]]),
            "f": int(row[col["Поверх"]]),
            "t": kind,
            "n": number,
            "paid": stairs + barriers + rest,
        }
        units.append(u)

    def sort_key(u):
        m = re.match(r"\D*(\d+)", u["n"])
        return (u["b"], u["s"], u["f"], u["t"] != "apt", int(m.group(1)) if m else 0, u["n"])

    units.sort(key=sort_key)

    if args.as_of:
        as_of = args.as_of
    else:
        dates = []
        for name in PAYMENT_SHEETS:
            if name in wb.sheetnames:
                for r in wb[name].iter_rows(min_row=2, values_only=True):
                    if isinstance(r[1], dt.datetime):
                        dates.append(r[1])
        as_of = max(dates).date().isoformat()

    out = {
        "asOf": as_of,
        "fullAmount": 8000,
        "addresses": ADDRESS,
        "byInitiative": totals,
        "units": units,
    }
    path = pathlib.Path(args.out)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(units)} units, total {sum(u['paid'] for u in units)} UAH, as of {as_of} -> {path}", file=sys.stderr)


if __name__ == "__main__":
    main()
